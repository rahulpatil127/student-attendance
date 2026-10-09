"""Session auth endpoints: CSRF bootstrap, login, logout, current user, admin user mgmt."""

from django.contrib.auth import authenticate, get_user_model, login, logout
from django.core.exceptions import ValidationError as DjangoValidationError
from django.middleware.csrf import get_token
from django.utils.decorators import method_decorator
from django.views.decorators.csrf import ensure_csrf_cookie
from drf_spectacular.utils import extend_schema
from rest_framework import status, viewsets
from rest_framework.exceptions import ValidationError as DRFValidationError
from rest_framework.permissions import AllowAny, IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from .permissions import IsAdmin
from .serializers import AdminUserSerializer, LoginSerializer, MeSerializer

User = get_user_model()


class CsrfView(APIView):
    permission_classes = [AllowAny]
    serializer_class = LoginSerializer

    @extend_schema(responses={200: {"type": "object"}})
    @method_decorator(ensure_csrf_cookie)
    def get(self, request):
        return Response({"detail": "CSRF cookie set", "csrfToken": get_token(request)})


class LoginView(APIView):
    permission_classes = [AllowAny]
    serializer_class = LoginSerializer

    @extend_schema(request=LoginSerializer, responses=MeSerializer)
    def post(self, request):
        serializer = LoginSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        username = serializer.validated_data["username"]
        password = serializer.validated_data["password"]
        user = authenticate(request, username=username, password=password)
        # Generic error: do not reveal whether username or password was wrong.
        if user is None or not user.is_active:
            return Response({"detail": "Invalid credentials."}, status=status.HTTP_400_BAD_REQUEST)
        login(request, user)  # rotates session key
        return Response(MeSerializer(user).data)


class LogoutView(APIView):
    permission_classes = [IsAuthenticated]
    serializer_class = MeSerializer

    @extend_schema(responses={200: {"type": "object"}})
    def post(self, request):
        logout(request)
        return Response({"detail": "Logged out."})


class MeView(APIView):
    permission_classes = [IsAuthenticated]
    serializer_class = MeSerializer

    @extend_schema(responses=MeSerializer)
    def get(self, request):
        return Response(MeSerializer(request.user).data)


class PasswordChangeView(APIView):
    """POST /auth/password/change/ — self-service password update."""

    permission_classes = [IsAuthenticated]

    @extend_schema(
        request={"type": "object"},
        responses={200: {"type": "object"}},
    )
    def post(self, request):
        from django.contrib.auth.password_validation import validate_password
        from django.core.exceptions import ValidationError as DjangoValidationError
        from rest_framework import serializers

        old = request.data.get("old_password", "")
        new = request.data.get("new_password", "")
        if not old or not new:
            raise serializers.ValidationError({"detail": "Both old_password and new_password are required."})
        user = request.user
        if not user.check_password(old):
            raise serializers.ValidationError({"old_password": "Current password is incorrect."})
        try:
            validate_password(new, user)
        except DjangoValidationError as exc:
            raise serializers.ValidationError({"new_password": list(exc.messages)}) from exc
        user.set_password(new)
        user.save(update_fields=["password"])
        return Response({"detail": "Password updated."})


class AdminUserViewSet(viewsets.ModelViewSet):
    queryset = User.objects.select_related("student_profile", "teacher_profile").prefetch_related(
        "enrollments__classroom", "enrollments__classroom__academic_year"
    ).all().order_by("username")
    serializer_class = AdminUserSerializer
    permission_classes = [IsAdmin]

    def get_queryset(self):
        from django.db.models import Q

        qs = super().get_queryset()
        p = self.request.query_params
        if p.get("role"):
            qs = qs.filter(role=p.get("role"))
        if p.get("is_active") in ("true", "True", "1"):
            qs = qs.filter(is_active=True)
        elif p.get("is_active") in ("false", "False", "0"):
            qs = qs.filter(is_active=False)
        if p.get("classroom"):
            qs = qs.filter(
                enrollments__classroom_id=p.get("classroom"),
                enrollments__status="ACTIVE",
            ).distinct()
        if p.get("search"):
            q = p.get("search")
            qs = qs.filter(
                Q(username__icontains=q)
                | Q(email__icontains=q)
                | Q(first_name__icontains=q)
                | Q(last_name__icontains=q)
            )
        return qs

    def perform_create(self, serializer):
        user = serializer.save()
        classroom_id = self.request.data.get("classroom")
        if not classroom_id:
            return
        if user.role != "STUDENT":
            raise DRFValidationError({"classroom": "Only students can be enrolled at creation."})
        from apps.academics.models import Classroom, Enrollment
        from apps.academics.views import ensure_student_profile

        try:
            classroom = Classroom.objects.get(pk=classroom_id)
        except (Classroom.DoesNotExist, ValueError, TypeError):
            raise DRFValidationError({"classroom": "Classroom not found."}) from None
        ensure_student_profile(user)
        try:
            enrollment = Enrollment(
                student=user,
                classroom=classroom,
                academic_year=classroom.academic_year,
                status="ACTIVE",
            )
            enrollment.full_clean()
            enrollment.save()
        except DjangoValidationError as exc:
            raise DRFValidationError(
                exc.message_dict if hasattr(exc, "message_dict") else str(exc)
            ) from exc


def _unique_username(base, taken):
    """Free username for bulk import: base, base2, base3… Returns (name, renamed)."""
    base = (base or "").strip()
    if base and base not in taken and not User.objects.filter(username=base).exists():
        return base, False
    i = 2
    while i < 1000:
        cand = f"{base}{i}"
        if cand not in taken and not User.objects.filter(username=cand).exists():
            return cand, True
        i += 1
    return None, False


def _normalize_header(name):
    return str(name or "").strip().lower().replace(" ", "_").replace("-", "_")


def _parse_bulk_rows(filename, content):
    """Return list of {row, username, firstname, lastname, email} or raise ValidationError."""
    import csv
    import io

    from rest_framework import serializers as drf_serializers

    name = filename.lower()
    if name.endswith(".csv"):
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        headers = [_normalize_header(h) for h in (reader.fieldnames or [])]
        rows = []
        for i, raw in enumerate(reader, start=2):
            item = {_normalize_header(k): (v or "").strip() for k, v in raw.items()}
            item["row"] = i
            rows.append(item)
    elif name.endswith((".xlsx", ".xls")):
        from openpyxl import load_workbook

        wb = load_workbook(filename=io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        grid = list(ws.iter_rows(values_only=True))
        if not grid:
            raise drf_serializers.ValidationError({"file": "Empty spreadsheet."})
        headers = [_normalize_header(h) for h in grid[0]]
        rows = []
        for i, values in enumerate(grid[1:], start=2):
            item = {
                headers[j]: ("" if v is None else str(v).strip())
                for j, v in enumerate(values)
                if j < len(headers)
            }
            item["row"] = i
            rows.append(item)
    else:
        raise drf_serializers.ValidationError(
            {"file": "Upload .xlsx or .csv (columns: username, firstname, lastname, email)."}
        )
    # unify header aliases
    out = []
    for item in rows:
        if not any(item.get(k) for k in ("username", "firstname", "lastname", "email")):
            continue  # skip blank lines
        out.append(
            {
                "row": item["row"],
                "username": item.get("username", ""),
                "firstname": item.get("firstname", "") or item.get("first_name", ""),
                "lastname": item.get("lastname", "") or item.get("last_name", ""),
                "email": item.get("email", ""),
            }
        )
    return out


class BulkUserUploadView(APIView):
    """POST /admin/users/bulk/ — create many STUDENT accounts from .xlsx/.csv.

    Columns: username, firstname, lastname (optional), email.
    Password defaults to the firstname; student numbers auto-generate.
    Taken usernames auto-resolve (sara → sara2) and are reported.
    Optional `classroom` form field enrolls every created student into it.
    Returns {created: [...], errors: [{row, username, reason}]}.
    """

    permission_classes = [IsAdmin]

    @extend_schema(request={"type": "object"}, responses={"type": "object"})
    def post(self, request):
        from django.core.exceptions import ValidationError as DjangoValidationError
        from rest_framework import serializers as drf_serializers

        upload = request.FILES.get("file")
        if upload is None:
            raise drf_serializers.ValidationError({"file": "Attach a .xlsx or .csv file."})
        if upload.size > 2 * 1024 * 1024:
            raise drf_serializers.ValidationError({"file": "File too large (max 2 MB)."})
        try:
            rows = _parse_bulk_rows(upload.name, upload.read())
        except drf_serializers.ValidationError:
            raise
        except Exception as exc:
            raise drf_serializers.ValidationError({"file": f"Could not read file: {exc}"}) from exc
        if len(rows) > 500:
            raise drf_serializers.ValidationError({"file": "Max 500 rows per upload."})

        from apps.academics.models import Classroom, Enrollment, StudentProfile
        from apps.academics.views import next_prefixed_number

        target_room = None
        if request.data.get("classroom"):
            try:
                target_room = Classroom.objects.get(pk=request.data.get("classroom"))
            except (Classroom.DoesNotExist, ValueError, TypeError):
                raise drf_serializers.ValidationError({"classroom": "Classroom not found."}) from None

        created, errors, seen_usernames, seen_emails = [], [], set(), set()
        for item in rows:
            username, email = item["username"], item["email"]
            firstname = item["firstname"]
            if not username or not email or not firstname:
                errors.append({"row": item["row"], "username": username, "reason": "username, firstname and email are required."})
                continue
            if email.lower() in seen_emails:
                errors.append({"row": item["row"], "username": username, "reason": "duplicate row in file."})
                continue
            try:
                if User.objects.filter(email__iexact=email).exists():
                    raise DjangoValidationError(f"Email {email} already exists.")
                final_name, renamed = _unique_username(username, seen_usernames)
                if final_name is None:
                    raise DjangoValidationError(f"Username {username} is taken and no variant is free.")
                seen_usernames.add(final_name)
                seen_emails.add(email.lower())
                user = User(
                    username=final_name,
                    email=email,
                    first_name=firstname,
                    last_name=item["lastname"],
                    role="STUDENT",
                )
                user.set_password(firstname)
                user.full_clean(exclude=["password"])
                user.save()
                profile, _ = StudentProfile.objects.get_or_create(
                    user=user,
                    defaults={"student_number": next_prefixed_number(StudentProfile, "student_number", "S")},
                )
                entry = {"username": final_name, "student_number": profile.student_number}
                if renamed:
                    entry["renamed_from"] = username
                if target_room is not None:
                    try:
                        enrollment = Enrollment(
                            student=user,
                            classroom=target_room,
                            academic_year=target_room.academic_year,
                            status="ACTIVE",
                        )
                        enrollment.full_clean()
                        enrollment.save()
                        entry["enrolled_in"] = str(target_room)
                    except DjangoValidationError as exc:
                        entry["enroll_error"] = "; ".join(exc.messages)
                created.append(entry)
            except DjangoValidationError as exc:
                errors.append({"row": item["row"], "username": username, "reason": "; ".join(exc.messages)})
        return Response({"created": created, "errors": errors, "total": len(rows)})
